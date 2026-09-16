"""Build the small, human-readable XLSX template for performance imports."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.importers.performance_record_file_parser import TEMPLATE_COLUMNS


def build_performance_record_template() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Performance Records"
    worksheet.append(list(TEMPLATE_COLUMNS))
    worksheet.append([
        "2026-09-01T00:00:00+00:00",
        "2026-09-02T00:00:00+00:00",
        1000,
        50,
        5,
        100.00,
        150.00,
        "示例数据，请替换或删除",
        None,
        None,
        None,
        None,
    ])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:L2"
    worksheet.sheet_view.showGridLines = False
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[2]:
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 24
    widths = {
        "A": 28, "B": 28, "C": 14, "D": 12, "E": 14,
        "F": 12, "G": 12, "H": 28, "I": 18, "J": 20,
        "K": 20, "L": 16,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet["F2"].number_format = "0.00"
    worksheet["G2"].number_format = "0.00"

    instructions = workbook.create_sheet("Instructions")
    instructions.sheet_view.showGridLines = False
    instructions.append(["字段", "说明"])
    instructions.append(["period_start / period_end", "ISO 8601，例如 2026-09-01T00:00:00+00:00"])
    instructions.append(["spend / revenue", "十进制金额，例如 100.00"])
    instructions.append(["optional relation IDs", "没有关联对象时留空"])
    instructions.append(["derived metrics", "CTR、Conversion Rate、ROI 由系统计算，不要添加这些列"])
    instructions.column_dimensions["A"].width = 30
    instructions.column_dimensions["B"].width = 72
    for cell in instructions[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in instructions.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
