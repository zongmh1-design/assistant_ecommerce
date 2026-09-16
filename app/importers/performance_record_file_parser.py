"""Parse CSV/XLSX cells into typed PerformanceRecord input values only."""

import csv
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.exceptions import AppError


MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_DATA_ROWS = 1000

REQUIRED_COLUMNS = (
    "period_start", "period_end", "impressions", "clicks",
    "conversions", "spend", "revenue",
)
OPTIONAL_COLUMNS = (
    "notes", "creative_plan_id", "generated_asset_id",
    "promotion_link_id", "experiment_id",
)
TEMPLATE_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
INTEGER_COLUMNS = {"impressions", "clicks", "conversions"}
OPTIONAL_ID_COLUMNS = {
    "creative_plan_id", "generated_asset_id",
    "promotion_link_id", "experiment_id",
}
DECIMAL_COLUMNS = {"spend", "revenue"}
DATETIME_COLUMNS = {"period_start", "period_end"}


@dataclass(frozen=True)
class ParseIssue:
    field: str | None
    error_code: str
    message: str


@dataclass
class ParsedPerformanceRow:
    row_number: int
    data: dict[str, object] = field(default_factory=dict)
    errors: list[ParseIssue] = field(default_factory=list)


class PerformanceRecordFileParser:
    def parse(self, filename: str | None, content: bytes) -> list[ParsedPerformanceRow]:
        if len(content) > MAX_IMPORT_FILE_BYTES:
            raise AppError(
                status_code=413, code="import_file_too_large",
                message="Import file exceeds the 5 MB limit",
            )
        suffix = Path(filename or "").suffix.lower()
        if suffix == ".csv":
            headers, rows = self._read_csv(content)
        elif suffix == ".xlsx":
            headers, rows = self._read_xlsx(content)
        else:
            raise _file_error(
                "unsupported_import_file_type", "Only .csv and .xlsx files are supported"
            )
        self._validate_headers(headers)
        parsed: list[ParsedPerformanceRow] = []
        for row_number, values in rows:
            if _is_empty_row(values):
                continue
            parsed.append(self._parse_row(row_number, headers, values))
            if len(parsed) > MAX_IMPORT_DATA_ROWS:
                raise AppError(
                    status_code=413, code="import_row_limit_exceeded",
                    message="Import file exceeds the 1000 data row limit",
                )
        return parsed

    def _read_csv(
        self, content: bytes
    ) -> tuple[list[str], list[tuple[int, list[object]]]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise _file_error(
                "invalid_csv_encoding", "CSV must use UTF-8 or UTF-8 with BOM"
            ) from exc
        reader = csv.reader(StringIO(text, newline=""))
        try:
            raw_headers = next(reader)
        except StopIteration as exc:
            raise _file_error("missing_header", "Import file does not contain a header row") from exc
        headers = [_header_value(value) for value in raw_headers]
        rows: list[tuple[int, list[object]]] = []
        data_row_count = 0
        for number, row in enumerate(reader, start=2):
            values = list(row)
            rows.append((number, values))
            if not _is_empty_row(values):
                data_row_count += 1
                _enforce_row_limit(data_row_count)
        return headers, rows

    def _read_xlsx(
        self, content: bytes
    ) -> tuple[list[str], list[tuple[int, list[object]]]]:
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=True
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise _file_error("invalid_xlsx", "XLSX file could not be read") from exc
        try:
            worksheet = workbook.worksheets[0]
            iterator = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(iterator)
            except StopIteration as exc:
                raise _file_error(
                    "missing_header", "Import file does not contain a header row"
                ) from exc
            headers = [_header_value(value) for value in raw_headers]
            rows: list[tuple[int, list[object]]] = []
            data_row_count = 0
            for number, row in enumerate(iterator, start=2):
                values = list(row)
                rows.append((number, values))
                if not _is_empty_row(values):
                    data_row_count += 1
                    _enforce_row_limit(data_row_count)
            return headers, rows
        finally:
            workbook.close()

    def _validate_headers(self, headers: list[str]) -> None:
        if not headers or all(not item for item in headers):
            raise _file_error("missing_header", "Import file does not contain a header row")
        duplicates = sorted({item for item in headers if item and headers.count(item) > 1})
        if duplicates:
            raise _file_error(
                "duplicate_columns", f"Duplicate columns: {', '.join(duplicates)}"
            )
        missing = sorted(set(REQUIRED_COLUMNS) - set(headers))
        if missing:
            raise _file_error(
                "missing_columns", f"Missing required columns: {', '.join(missing)}"
            )
        unknown = sorted({item for item in headers if item not in TEMPLATE_COLUMNS})
        if unknown:
            raise _file_error(
                "unknown_columns", f"Unknown columns: {', '.join(unknown)}"
            )

    def _parse_row(
        self, row_number: int, headers: list[str], values: list[object]
    ) -> ParsedPerformanceRow:
        row = ParsedPerformanceRow(row_number=row_number)
        padded = values + [None] * max(0, len(headers) - len(values))
        if len(values) > len(headers) and not _is_empty_row(values[len(headers):]):
            row.errors.append(
                ParseIssue(
                    field=None,
                    error_code="extra_cells",
                    message="Row contains values beyond the declared columns",
                )
            )
        for name, value in zip(headers, padded, strict=False):
            try:
                if name in DATETIME_COLUMNS:
                    row.data[name] = _parse_datetime(value)
                elif name in INTEGER_COLUMNS:
                    row.data[name] = _parse_integer(value, required=True)
                elif name in DECIMAL_COLUMNS:
                    row.data[name] = _parse_decimal(value)
                elif name in OPTIONAL_ID_COLUMNS:
                    row.data[name] = _parse_optional_id(value)
                elif name == "notes":
                    row.data[name] = None if _is_blank(value) else str(value).strip()
            except ValueError as exc:
                row.errors.append(
                    ParseIssue(
                        field=name,
                        error_code=_parse_error_code(name),
                        message=str(exc),
                    )
                )
        return row


def _parse_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str) and value.strip():
        text = value.strip()
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("must be an ISO 8601 datetime") from exc
    else:
        raise ValueError("is required and must be an ISO 8601 datetime")
    if parsed.tzinfo is None:
        # Excel datetime cells do not retain timezone metadata; interpret them as UTC.
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _parse_integer(value: object, *, required: bool) -> int | None:
    if _is_blank(value):
        if required:
            raise ValueError("is required and must be an integer")
        return None
    if isinstance(value, bool):
        raise ValueError("must be an integer")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("must be an integer") from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError("must be an integer")
    return int(number)


def _parse_optional_id(value: object) -> int | None:
    parsed = _parse_integer(value, required=False)
    if parsed is not None and parsed <= 0:
        raise ValueError("must be a positive integer when provided")
    return parsed


def _parse_decimal(value: object) -> Decimal:
    if _is_blank(value):
        raise ValueError("is required and must be a decimal number")
    if isinstance(value, bool):
        raise ValueError("must be a decimal number")
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("must be a decimal number") from exc
    if not parsed.is_finite():
        raise ValueError("must be a finite decimal number")
    return parsed


def _header_value(value: object) -> str:
    return "" if value is None else str(value).strip()


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_empty_row(values: list[object]) -> bool:
    return all(_is_blank(value) for value in values)


def _parse_error_code(field_name: str) -> str:
    if field_name in DATETIME_COLUMNS:
        return "invalid_datetime"
    if field_name in INTEGER_COLUMNS:
        return "invalid_integer"
    if field_name in OPTIONAL_ID_COLUMNS:
        return "invalid_optional_id"
    if field_name in DECIMAL_COLUMNS:
        return "invalid_decimal"
    return "invalid_value"


def _enforce_row_limit(data_row_count: int) -> None:
    if data_row_count > MAX_IMPORT_DATA_ROWS:
        raise AppError(
            status_code=413, code="import_row_limit_exceeded",
            message="Import file exceeds the 1000 data row limit",
        )


def _file_error(code: str, message: str) -> AppError:
    return AppError(status_code=422, code=code, message=message)
