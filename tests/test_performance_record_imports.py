"""CSV/XLSX template, stateless preview and partial import tests."""

import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.importers.performance_record_file_parser import (
    MAX_IMPORT_DATA_ROWS,
    MAX_IMPORT_FILE_BYTES,
    TEMPLATE_COLUMNS,
)
from app.models.ad_experiment import AdExperiment, AdExperimentStatus
from app.models.ad_recommendation import (
    AdRecommendation,
    AdRecommendationConfirmStatus,
)
from app.models.creative_plan import (
    CreativePlan,
    CreativePlanStatus,
    CreativePlanType,
)
from app.models.performance_record import PerformanceRecord
from app.models.product import Product, ProductStatus
from app.models.store import Platform, Store
from app.models.user import UserRole


def add_product(db: Session, suffix: str = "a") -> Product:
    store = Store(store_name=f"Import Store {suffix}", platform=Platform.JD)
    db.add(store)
    db.flush()
    product = Product(
        store_id=store.id,
        name=f"Import Product {suffix}",
        platform=Platform.JD,
        category="家居",
        price=Decimal("199.90"),
        status=ProductStatus.ACTIVE,
    )
    db.add(product)
    db.commit()
    db.refresh(product)
    return product


def add_plan(db: Session, product: Product, suffix: str = "a") -> CreativePlan:
    plan = CreativePlan(
        product_id=product.id,
        plan_type=CreativePlanType.MAIN_IMAGE,
        title=f"Import Plan {suffix}",
        content_json={"core_copy": "fixture"},
        rationale_text="fixture",
        status=CreativePlanStatus.DRAFT,
        input_context_json={"product": {"name": product.name}},
    )
    db.add(plan)
    db.commit()
    db.refresh(plan)
    return plan


def add_experiment(
    db: Session,
    product: Product,
    status: AdExperimentStatus,
    suffix: str = "a",
) -> AdExperiment:
    recommendation = AdRecommendation(
        product_id=product.id,
        summary_text="fixture",
        objective_text="fixture",
        audience_segments_json=[],
        budget_plan_json={
            "total_budget": "100.00",
            "currency": "CNY",
            "allocation": [],
            "rationale": "fixture",
        },
        creative_tests_json=[],
        bid_strategy_json={
            "strategy_name": "manual",
            "rationale": "fixture",
            "constraints": [],
        },
        risk_controls_json=[],
        next_steps_json=[],
        confirm_status=AdRecommendationConfirmStatus.CONFIRMED,
        input_context_json={},
    )
    db.add(recommendation)
    db.flush()
    experiment = AdExperiment(
        product_id=product.id,
        ad_recommendation_id=recommendation.id,
        experiment_name=f"Import Experiment {suffix}",
        target_text="fixture",
        audience_text="fixture",
        budget_amount=Decimal("100.00"),
        success_metric_text="fixture",
        hypothesis_text="fixture",
        experiment_status=status,
        input_context_json={},
    )
    db.add(experiment)
    db.commit()
    db.refresh(experiment)
    return experiment


def valid_row(**overrides):
    row = {
        "period_start": "2026-09-01T00:00:00+00:00",
        "period_end": "2026-09-02T00:00:00+00:00",
        "impressions": "1000",
        "clicks": "50",
        "conversions": "5",
        "spend": "100.00",
        "revenue": "150.00",
        "notes": "manual import",
        "creative_plan_id": "",
        "generated_asset_id": "",
        "promotion_link_id": "",
        "experiment_id": "",
    }
    row.update(overrides)
    return row


def csv_file(rows, headers=TEMPLATE_COLUMNS, *, bom=False) -> bytes:
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=list(headers), extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        if row is None:
            output.write("," * (len(headers) - 1) + "\r\n")
        else:
            writer.writerow(row)
    encoded = output.getvalue().encode("utf-8")
    return (b"\xef\xbb\xbf" + encoded) if bom else encoded


def xlsx_file(
    rows: list[list[object]], headers=TEMPLATE_COLUMNS, second_sheet=None
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(row)
    if second_sheet is not None:
        other = workbook.create_sheet("Ignored")
        for row in second_sheet:
            other.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def upload(client, path, headers, filename, content, content_type):
    return client.post(
        path,
        files={"file": (filename, content, content_type)},
        headers=headers,
    )


def test_template_download_has_exact_headers_example_and_no_derived_columns(
    client, auth_headers_factory
):
    response = client.get(
        "/api/v1/workspace/templates/performance-records",
        headers=auth_headers_factory(UserRole.VIEWER),
    )
    assert response.status_code == 200
    assert "performance-records-template.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    worksheet = workbook.worksheets[0]
    assert [cell.value for cell in worksheet[1]] == list(TEMPLATE_COLUMNS)
    assert worksheet["A2"].value == "2026-09-01T00:00:00+00:00"
    assert worksheet["F2"].value == 100
    assert not {"ctr", "conversion_rate", "roi"} & set(TEMPLATE_COLUMNS)
    assert worksheet.freeze_panes == "A2"
    workbook.close()


@pytest.mark.parametrize("bom", [False, True])
def test_csv_preview_accepts_utf8_and_bom_without_writing(
    client, db_session, auth_headers_factory, bom
):
    product = add_product(db_session)
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import/preview",
        auth_headers_factory(UserRole.OPERATOR),
        "records.csv",
        csv_file([valid_row()], bom=bom),
        "text/csv",
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total_rows"] == 1
    assert body["valid_rows"] == 1
    assert body["rows"][0]["row_number"] == 2
    assert body["rows"][0]["calculated_metrics"] == {
        "ctr": "0.050000",
        "conversion_rate": "0.100000",
        "roi": "0.500000",
    }
    assert db_session.scalar(
        select(func.count()).select_from(PerformanceRecord)
    ) == 0


def test_preview_rejects_missing_and_unknown_headers(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    headers = auth_headers_factory(UserRole.ADMIN)
    path = f"/api/v1/products/{product.id}/performance-records/import/preview"
    missing = tuple(name for name in TEMPLATE_COLUMNS if name != "revenue")
    response = upload(
        client, path, headers, "missing.csv",
        csv_file([valid_row()], missing), "text/csv",
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "missing_columns"
    unknown = TEMPLATE_COLUMNS + ("click_rate",)
    response = upload(
        client, path, headers, "unknown.csv",
        csv_file([valid_row()], unknown), "text/csv",
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "unknown_columns"


def test_empty_rows_are_skipped_and_header_only_is_valid(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    path = f"/api/v1/products/{product.id}/performance-records/import/preview"
    headers = auth_headers_factory(UserRole.ADMIN)
    response = upload(
        client, path, headers, "blank.csv",
        csv_file([None, valid_row(), None]), "text/csv",
    )
    assert response.status_code == 200
    assert response.json()["total_rows"] == 1
    assert response.json()["rows"][0]["row_number"] == 3
    empty = upload(
        client, path, headers, "empty.csv",
        csv_file([]), "text/csv",
    )
    assert empty.status_code == 200
    assert empty.json()["total_rows"] == 0


@pytest.mark.parametrize(
    ("field", "value", "code"),
    [
        ("impressions", "10.5", "invalid_integer"),
        ("spend", "money", "invalid_decimal"),
        ("period_start", "2026/09/01", "invalid_datetime"),
    ],
)
def test_csv_basic_type_errors_have_field_and_row(
    client, db_session, auth_headers_factory, field, value, code
):
    product = add_product(db_session)
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import/preview",
        auth_headers_factory(UserRole.ADMIN),
        "invalid.csv",
        csv_file([valid_row(**{field: value})]),
        "text/csv",
    )
    assert response.status_code == 200
    error = response.json()["rows"][0]["errors"][0]
    assert error["field"] == field
    assert error["error_code"] == code
    assert response.json()["rows"][0]["row_number"] == 2


@pytest.mark.parametrize(
    ("overrides", "code"),
    [
        ({"impressions": "10", "clicks": "11"}, "clicks_exceed_impressions"),
        ({"clicks": "2", "conversions": "3"}, "conversions_exceed_clicks"),
    ],
)
def test_preview_reuses_phase11a_metric_relationship_rules(
    client, db_session, auth_headers_factory, overrides, code
):
    product = add_product(db_session)
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import/preview",
        auth_headers_factory(UserRole.ADMIN),
        "business.csv",
        csv_file([valid_row(**overrides)]),
        "text/csv",
    )
    body = response.json()
    assert body["invalid_rows"] == 1
    assert body["rows"][0]["errors"][0]["error_code"] == code


def test_preview_optional_relation_valid_and_cross_product_invalid(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session, "a")
    other = add_product(db_session, "b")
    plan = add_plan(db_session, product)
    foreign_plan = add_plan(db_session, other, "b")
    path = f"/api/v1/products/{product.id}/performance-records/import/preview"
    headers = auth_headers_factory(UserRole.ADMIN)
    response = upload(
        client, path, headers, "relations.csv",
        csv_file([
            valid_row(creative_plan_id=str(plan.id)),
            valid_row(creative_plan_id=str(foreign_plan.id)),
        ]), "text/csv",
    )
    body = response.json()
    assert body["valid_rows"] == 1 and body["invalid_rows"] == 1
    assert body["rows"][1]["errors"][0]["error_code"] == "creative_plan_not_found"


@pytest.mark.parametrize(
    ("status", "valid"),
    [
        (AdExperimentStatus.DRAFT, False),
        (AdExperimentStatus.CONFIRMED, False),
        (AdExperimentStatus.RUNNING, True),
        (AdExperimentStatus.FINISHED, True),
    ],
)
def test_preview_reuses_experiment_status_rule(
    client, db_session, auth_headers_factory, status, valid
):
    product = add_product(db_session)
    experiment = add_experiment(db_session, product, status)
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import/preview",
        auth_headers_factory(UserRole.OPERATOR),
        "experiment.csv",
        csv_file([valid_row(experiment_id=str(experiment.id))]),
        "text/csv",
    )
    body = response.json()
    assert body["valid_rows"] == int(valid)
    assert body["invalid_rows"] == int(not valid)


def test_xlsx_preview_accepts_datetime_decimal_blank_ids_and_reads_first_sheet_only(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    row = [
        datetime(2026, 9, 1), datetime(2026, 9, 2),
        1000.0, 50.0, 5.0, 100.25, 150.75, "xlsx",
        None, None, None, None,
    ]
    content = xlsx_file(
        [row],
        second_sheet=[["bad", "headers"], ["must", "be ignored"]],
    )
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import/preview",
        auth_headers_factory(UserRole.ADMIN),
        "records.xlsx",
        content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert response.status_code == 200
    body = response.json()
    assert body["valid_rows"] == 1
    normalized = body["rows"][0]["normalized_data"]
    assert normalized["period_start"] == "2026-09-01T00:00:00Z"
    assert normalized["spend"] == "100.25"
    assert normalized["creative_plan_id"] is None


def test_xlsx_error_preserves_original_row_number(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    blank = [None] * len(TEMPLATE_COLUMNS)
    invalid = [
        "not-a-date", "2026-09-02T00:00:00+00:00",
        100, 10, 1, 10, 20, None, None, None, None, None,
    ]
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import/preview",
        auth_headers_factory(UserRole.ADMIN),
        "rows.xlsx",
        xlsx_file([blank, invalid]),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert response.json()["rows"][0]["row_number"] == 3


def test_import_all_valid_and_duplicate_file_is_allowed(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    content = csv_file([valid_row(), valid_row(notes="second")])
    path = f"/api/v1/products/{product.id}/performance-records/import"
    headers = auth_headers_factory(UserRole.ADMIN)
    first = upload(client, path, headers, "records.csv", content, "text/csv")
    second = upload(client, path, headers, "records.csv", content, "text/csv")
    assert first.status_code == 200 and second.status_code == 200
    assert first.json()["success_count"] == 2
    assert all(item["performance_record_id"] > 0 for item in first.json()["successes"])
    assert db_session.scalar(
        select(func.count()).select_from(PerformanceRecord)
    ) == 4


def test_import_revalidates_after_preview_when_experiment_status_changes(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    experiment = add_experiment(
        db_session, product, AdExperimentStatus.RUNNING
    )
    content = csv_file([valid_row(experiment_id=str(experiment.id))])
    base = f"/api/v1/products/{product.id}/performance-records/import"
    headers = auth_headers_factory(UserRole.ADMIN)
    preview = upload(
        client, base + "/preview", headers, "experiment.csv", content, "text/csv"
    )
    assert preview.json()["valid_rows"] == 1
    experiment.experiment_status = AdExperimentStatus.CANCELLED
    db_session.commit()
    result = upload(
        client, base, headers, "experiment.csv", content, "text/csv"
    )
    assert result.json()["success_count"] == 0
    assert result.json()["failures"][0]["errors"][0]["error_code"] == "ad_experiment_not_started"


def test_import_partially_succeeds_and_calculates_metrics(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    content = csv_file([
        valid_row(),
        valid_row(impressions="10", clicks="11"),
        valid_row(spend="0.00", revenue="10.00", notes="zero spend"),
    ])
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import",
        auth_headers_factory(UserRole.OPERATOR),
        "partial.csv",
        content,
        "text/csv",
    )
    body = response.json()
    assert body["total_rows"] == 3
    assert body["success_count"] == 2
    assert body["failure_count"] == 1
    assert body["failures"][0]["row_number"] == 3
    records = list(db_session.scalars(select(PerformanceRecord).order_by(PerformanceRecord.id)))
    assert len(records) == 2
    assert records[0].ctr == Decimal("0.050000")
    assert records[0].conversion_rate == Decimal("0.100000")
    assert records[0].roi == Decimal("0.500000")
    assert records[1].roi is None


def test_import_all_invalid_writes_nothing(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    response = upload(
        client,
        f"/api/v1/products/{product.id}/performance-records/import",
        auth_headers_factory(UserRole.ADMIN),
        "invalid.csv",
        csv_file([valid_row(clicks="bad"), valid_row(spend="bad")]),
        "text/csv",
    )
    assert response.json()["success_count"] == 0
    assert response.json()["failure_count"] == 2
    assert db_session.scalar(
        select(func.count()).select_from(PerformanceRecord)
    ) == 0


def test_viewer_cannot_preview_or_import(client, db_session, auth_headers_factory):
    product = add_product(db_session)
    headers = auth_headers_factory(UserRole.VIEWER)
    content = csv_file([valid_row()])
    base = f"/api/v1/products/{product.id}/performance-records/import"
    assert upload(client, base + "/preview", headers, "a.csv", content, "text/csv").status_code == 403
    assert upload(client, base, headers, "a.csv", content, "text/csv").status_code == 403


def test_file_size_and_row_limits(client, db_session, auth_headers_factory):
    product = add_product(db_session)
    path = f"/api/v1/products/{product.id}/performance-records/import/preview"
    headers = auth_headers_factory(UserRole.ADMIN)
    too_large = upload(
        client, path, headers, "large.csv",
        b"x" * (MAX_IMPORT_FILE_BYTES + 1), "text/csv",
    )
    assert too_large.status_code == 413
    rows = [valid_row(notes=str(index)) for index in range(MAX_IMPORT_DATA_ROWS + 1)]
    too_many = upload(
        client, path, headers, "rows.csv", csv_file(rows), "text/csv",
    )
    assert too_many.status_code == 413
    assert too_many.json()["error"]["code"] == "import_row_limit_exceeded"


def test_empty_file_invalid_encoding_and_unsupported_type_are_rejected(
    client, db_session, auth_headers_factory
):
    product = add_product(db_session)
    path = f"/api/v1/products/{product.id}/performance-records/import/preview"
    headers = auth_headers_factory(UserRole.ADMIN)
    empty = upload(client, path, headers, "empty.csv", b"", "text/csv")
    assert empty.status_code == 422
    invalid_encoding = upload(
        client, path, headers, "gbk.csv", "经营数据".encode("gbk"), "text/csv"
    )
    assert invalid_encoding.status_code == 422
    unsupported = upload(
        client, path, headers, "data.xls", b"legacy", "application/vnd.ms-excel"
    )
    assert unsupported.status_code == 422
